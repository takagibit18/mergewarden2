#!/usr/bin/env python3
"""Build DOCX/Markdown from one source. Requires python-docx; fonts are not bundled."""
from pathlib import Path
import json
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
R=Path(__file__).resolve().parents[1]
D=json.loads((R/'docs/design-source.json').read_text(encoding='utf8'))
FONT='Noto Sans CJK SC'; NAVY='173149'; TEAL='167D8D'; GRAY='5D6B78'
def fmt(run,size=10.5,bold=False,color='263747'):
 run.font.name=FONT;run.font.size=Pt(size);run.font.bold=bold;run.font.color.rgb=RGBColor.from_string(color)
 rf=run._element.get_or_add_rPr().rFonts
 for k in ('ascii','hAnsi','eastAsia','cs'):rf.set(qn('w:'+k),FONT)
def p(doc,text,style=None,size=10.5,bold=False,color='263747'):
 x=doc.add_paragraph(style=style);x.paragraph_format.space_after=Pt(5);x.paragraph_format.line_spacing=1.25
 fmt(x.add_run(text),size,bold,color);return x
def shade(cell,color):
 sh=OxmlElement('w:shd');sh.set(qn('w:fill'),color);cell._tc.get_or_add_tcPr().append(sh)
def table(doc,headers,rows):
 t=doc.add_table(rows=1,cols=len(headers));t.autofit=False;t.alignment=WD_TABLE_ALIGNMENT.CENTER
 if len(headers)==2: widths=[3.44,3.44] if headers[0]=='明确方向' else [2.08,4.8]
 elif headers[0]=='ADR':widths=[.58,3.3,3.0]
 else:widths=[1.3,2.8,2.78]
 for i,w in enumerate(widths):t.columns[i].width=Inches(w)
 for i,h in enumerate(headers):
  c=t.rows[0].cells[i];c.width=Inches(widths[i]);shade(c,NAVY);x=c.paragraphs[0];x.paragraph_format.space_before=Pt(3);x.paragraph_format.space_after=Pt(3);fmt(x.add_run(h),9.5,True,'FFFFFF')
 t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
 for ri,row in enumerate(rows):
  cells=t.add_row().cells;cells[0]._tc.getparent().get_or_add_trPr().append(OxmlElement('w:cantSplit'))
  for i,text in enumerate(row):
   c=cells[i];c.width=Inches(widths[i]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER;shade(c,'F0F5F8' if ri%2==0 else 'FAFCFD')
   mar=OxmlElement('w:tcMar')
   for name,val in [('top','65'),('bottom','65'),('left','95'),('right','95')]:
    e=OxmlElement('w:'+name);e.set(qn('w:w'),val);e.set(qn('w:type'),'dxa');mar.append(e)
   c._tc.get_or_add_tcPr().append(mar);x=c.paragraphs[0];x.paragraph_format.space_after=Pt(1);x.paragraph_format.space_before=Pt(1);x.paragraph_format.line_spacing=1.15;fmt(x.add_run(str(text)),9.2,i==0,NAVY if i==0 else '263747')
 gap=doc.add_paragraph();gap.paragraph_format.space_after=Pt(0);gap.paragraph_format.space_before=Pt(0);gap.paragraph_format.line_spacing=Pt(3);fmt(gap.add_run(''),3)
def callout(doc,text):
 x=doc.add_paragraph();x.paragraph_format.space_before=Pt(5);x.paragraph_format.space_after=Pt(8);x.paragraph_format.left_indent=Inches(.12);x.paragraph_format.right_indent=Inches(.1);x.paragraph_format.line_spacing=1.24
 pr=x._p.get_or_add_pPr();sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'EAF4F5');pr.append(sh);bd=OxmlElement('w:pBdr');left=OxmlElement('w:left')
 for k,v in [('val','single'),('sz','18'),('color',TEAL),('space','8')]:left.set(qn('w:'+k),v)
 bd.append(left);pr.append(bd);fmt(x.add_run(text),10.3,True,NAVY)
def link(x,label,url):
 h=OxmlElement('w:hyperlink');h.set(qn('r:id'),x.part.relate_to(url,RT.HYPERLINK,is_external=True));r=OxmlElement('w:r');rp=OxmlElement('w:rPr')
 for tag,attr,value in [('color','val',TEAL),('u','val','single'),('sz','val','18')]:
  el=OxmlElement('w:'+tag);el.set(qn('w:'+attr),value);rp.append(el)
 r.append(rp);txt=OxmlElement('w:t');txt.text=label;r.append(txt);h.append(r);x._p.append(h)
def build():
 doc=Document();s=doc.sections[0];s.page_width=Inches(8.2677);s.page_height=Inches(11.6929);s.top_margin=Inches(.65);s.bottom_margin=Inches(.6);s.left_margin=Inches(.69);s.right_margin=Inches(.69);s.header_distance=Inches(.25);s.footer_distance=Inches(.25);s.different_first_page_header_footer=True
 for name in ('Normal','Title','Subtitle','Heading 1','Heading 2'):
  st=doc.styles[name];st.font.name=FONT;st._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),FONT)
 doc.styles['Normal'].font.size=Pt(10.5);doc.styles['Normal'].paragraph_format.widow_control=True
 for name in ('Heading 1','Heading 2'):doc.styles[name].paragraph_format.keep_with_next=True
 fmt(s.header.paragraphs[0].add_run('MERGEWARDEN 2.0  /  ENGINEERING DESIGN'),8,True,GRAY)
 ft=s.footer.paragraphs[0];ft.alignment=WD_ALIGN_PARAGRAPH.RIGHT;fmt(ft.add_run('设计基线 0.1  ·  2026-09-20     |     '),8,False,GRAY);fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');ft._p.append(fld)
 fmt(s.first_page_footer.paragraphs[0].add_run('ARCHITECTURE SCAFFOLD  ·  NOT A PRODUCTION RELEASE'),8,True,GRAY)
 doc.core_properties.title='MergeWarden 2.0 顶层设计与决策记录';doc.core_properties.author='MergeWarden Project';doc.core_properties.subject='Pi、CodeGraph、Finding 策略、Jev 与产品入口'
 md=['# MergeWarden 2.0 顶层设计与决策记录','',f"版本：{D['version']} | 日期：{D['date']}",'']
 for index,page in enumerate(D['pages']):
  if index:doc.add_page_break()
  x=p(doc,page['kicker'],size=8.5,bold=True,color=TEAL);x.paragraph_format.space_after=Pt(7)
  x=p(doc,page['title'],style='Title' if index==0 else 'Heading 1',size=35 if index==0 else 22,bold=True,color=NAVY);x.paragraph_format.space_after=Pt(10)
  md+=['## '+page['title'],'']
  for b in page['blocks']:
   k=b['kind']
   if k=='p':p(doc,b['text']);md += [b['text'],'']
   elif k=='h2':
    x=p(doc,b['text'],style='Heading 2',size=12.4,bold=True,color=NAVY);x.paragraph_format.space_before=Pt(7);x.paragraph_format.space_after=Pt(4);md+=['### '+b['text'],'']
   elif k=='subtitle':p(doc,b['text'],style='Subtitle',size=17,color=GRAY);md += [b['text'],'']
   elif k=='callout':callout(doc,b['text']);md+=['> '+b['text'],'']
   elif k=='table':
    table(doc,b['headers'],b['rows']);md+=['| '+' | '.join(b['headers'])+' |','| '+' | '.join('---' for _ in b['headers'])+' |'];md+=['| '+' | '.join(str(v).replace('|',' / ') for v in row)+' |' for row in b['rows']]+['']
  if index==len(D['pages'])-1:
   x=doc.add_paragraph();fmt(x.add_run('原始资料： '),9,True)
   for i,source in enumerate(D['sources']):
    if i:fmt(x.add_run('   '),9)
    link(x,source['id'],source['url'])
 out=R/'docs/MergeWarden2_Top_Level_Design.docx';doc.save(out);(R/'docs/ARCHITECTURE.md').write_text('\n'.join(md)+'\n',encoding='utf8');print(out)
if __name__=='__main__':build()
